# -*- coding: utf-8 -*-
"""
Automatização CCMSNS para o Conversor CCF / Mapeamentos CCM.

Funcionalidades:
- Outlook clássico no Windows;
- histórico 2026;
- processamento de novos emails desde 16/09/2026;
- despesas faturadas/conferidas;
- reclamações;
- gravação estruturada na pasta de rede;
- log de processamento em Excel.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import tempfile
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional, Tuple

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook


CCMSNS_SENDER = "ccmsns@spms.min-saude.pt"
TARGET_ATTACHMENT = "ULS Litoral Alentejano.zip"

PROCESS_FROM = datetime(2026, 9, 16, 0, 0, 0)
HISTORY_FROM = datetime(2026, 1, 1, 0, 0, 0)
HISTORY_UNTIL = datetime(2026, 9, 15, 23, 59, 59)

CONFERENCIA_ROOT = Path(
    r"G:\Administrativo\Servicos Financeiros\CONTABILIDADE - ULSLA"
    r"\Convencionados - conferência - CCF Maia\Conferencia"
)
LOG_YEAR = 2026
LOG_PATH = CONFERENCIA_ROOT / str(LOG_YEAR) / "Log_Processamento_CCF.xlsx"
OUTLOOK_SCAN_SUBFOLDERS = True

MONTHS_PT = {
    "janeiro": (1, "Janeiro"),
    "fevereiro": (2, "Fevereiro"),
    "marco": (3, "Março"),
    "abril": (4, "Abril"),
    "maio": (5, "Maio"),
    "junho": (6, "Junho"),
    "julho": (7, "Julho"),
    "agosto": (8, "Agosto"),
    "setembro": (9, "Setembro"),
    "outubro": (10, "Outubro"),
    "novembro": (11, "Novembro"),
    "dezembro": (12, "Dezembro"),
}

LOG_HEADERS = [
    "Data do registo",
    "Data do email",
    "Remetente",
    "Assunto",
    "MessageID",
    "Tipo",
    "Área",
    "Ano referência",
    "Mês referência",
    "Estado",
    "Anexo(s)",
    "Pasta destino",
    "Modo",
    "Ação efetuada",
    "Resultado",
    "Ficheiros gravados",
    "Convenções em falta",
    "CC fallback 9197",
    "Observações",
]


def normalize_text(value: str) -> str:
    value = str(value or "").strip().lower()
    value = unicodedata.normalize("NFD", value)
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", value).strip()


def normalize_filename(value: str) -> str:
    value = normalize_text(value)
    return re.sub(r"\s*\.\s*", ".", value)


def normalize_datetime(value) -> Optional[datetime]:
    if value is None:
        return None
    try:
        return datetime(
            value.year, value.month, value.day,
            value.hour, value.minute, value.second,
        )
    except Exception:
        return None


def classify_ccmsns_area(text: str) -> Tuple[Optional[str], str]:
    raw = str(text or "").strip()
    norm = normalize_text(raw)
    matches: list[str] = []

    def add(area: str) -> None:
        if area not in matches:
            matches.append(area)

    if "mcdt" in norm or "meios complementares de diagnostico" in norm:
        add("MCDTs")
    if "hemodial" in norm:
        add("Hemodialise")
    if "medicamento" in norm or "farmac" in norm:
        add("Medicamentos")
    if "proximidade" in norm:
        add("Proximidade")
    if ("tratamento" in norm and "term" in norm) or "termal" in norm:
        add("Tratamentos Termais")
    if "vacina" in norm:
        add("Vacinas")
    if "saude oral" in norm or "siso" in norm:
        add("Saúde Oral")
    if re.search(r"(^|\s)crd($|\s)", norm) or "cuidados respiratorios" in norm:
        add("CRD")

    return (matches[0], raw) if len(matches) == 1 else (None, raw)


def extract_month_year_from_text(text: str) -> Tuple[Optional[str], Optional[int]]:
    raw = str(text or "")
    norm = normalize_text(raw)
    found: list[tuple[str, int]] = []

    def add(month_name: str, year: int) -> None:
        pair = (month_name, int(year))
        if pair not in found:
            found.append(pair)

    for month_key, (_num, month_name) in MONTHS_PT.items():
        for match in re.finditer(
            rf"\b{re.escape(month_key)}\b\s*(?:de\s*)?(20\d{{2}})\b", norm
        ):
            add(month_name, int(match.group(1)))

    for match in re.finditer(r"\b(0?[1-9]|1[0-2])[./-](20\d{2})\b", norm):
        month_num, year = int(match.group(1)), int(match.group(2))
        for _key, (num, name) in MONTHS_PT.items():
            if num == month_num:
                add(name, year)
                break

    for match in re.finditer(r"\b(20\d{2})[./-](0?[1-9]|1[0-2])\b", norm):
        year, month_num = int(match.group(1)), int(match.group(2))
        for _key, (num, name) in MONTHS_PT.items():
            if num == month_num:
                add(name, year)
                break

    return found[0] if len(found) == 1 else (None, None)


def parse_ccmsns_subject(subject: str) -> Optional[dict]:
    raw = str(subject or "").strip()
    match = re.search(
        r"^\s*Despesa\s+(Faturada|Conferida)\s+de\s+(.+?)"
        r"\s*-\s*([A-Za-zÀ-ÿ]+)\s+de\s+(\d{4})\s*$",
        raw,
        flags=re.IGNORECASE,
    )
    if not match:
        return None

    tipo_raw, area_raw, mes_raw, ano_raw = match.groups()
    mes_norm = normalize_text(mes_raw)
    if mes_norm not in MONTHS_PT:
        return None

    _, mes_nome = MONTHS_PT[mes_norm]
    area, _ = classify_ccmsns_area(area_raw)
    tipo = "Faturada" if normalize_text(tipo_raw) == "faturada" else "Conferida"

    return {
        "tipo": tipo,
        "area": area,
        "area_original": area_raw.strip(),
        "ano": int(ano_raw),
        "mes": mes_nome,
        "estado": "Não conferido" if tipo == "Faturada" else "Conferido",
    }


def get_mail_body(mail) -> str:
    try:
        return str(getattr(mail, "Body", "") or "")
    except Exception:
        return ""


def is_reclamacao_email(subject: str, body: str = "") -> bool:
    return "reclamac" in normalize_text(f"{subject or ''}\n{body or ''}")


def parse_reclamacao_email(subject: str, body: str = "") -> Optional[dict]:
    if not is_reclamacao_email(subject, body):
        return None

    area, area_original = classify_ccmsns_area(subject)
    if area is None:
        area, area_original = classify_ccmsns_area(body)

    mes, ano = extract_month_year_from_text(subject)
    if mes is None or ano is None:
        mes, ano = extract_month_year_from_text(body)

    return {
        "tipo": "Reclamação",
        "area": area,
        "area_original": area_original,
        "ano": ano,
        "mes": mes,
        "estado": "Reclamações",
    }


def build_destination(parsed: dict) -> Optional[Path]:
    if not parsed or not parsed.get("area") or not parsed.get("ano") or not parsed.get("mes"):
        return None
    return (
        CONFERENCIA_ROOT
        / str(parsed["ano"])
        / parsed["mes"]
        / parsed["area"]
        / parsed["estado"]
    )


def get_outlook_namespace():
    if os.name != "nt":
        raise RuntimeError("A leitura automática do Outlook só funciona no Windows.")

    try:
        import pythoncom
        pythoncom.CoInitialize()
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "Falta instalar o módulo pywin32. Execute novamente INSTALAR_LOCAL.bat."
        ) from exc

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        return outlook.GetNamespace("MAPI")
    except Exception as exc:
        raise RuntimeError(
            "Não foi possível ligar ao Outlook clássico. "
            "Confirme que o Outlook está instalado e que existe um perfil configurado."
        ) from exc


def testar_ligacao_outlook() -> dict:
    try:
        namespace = get_outlook_namespace()
        inbox = namespace.GetDefaultFolder(6)
        return {
            "ok": True,
            "mensagem": f"Ligação ao Outlook concluída. Caixa de Entrada: {inbox.Name}",
        }
    except Exception as exc:
        return {"ok": False, "mensagem": str(exc)}


def get_sender_smtp(mail) -> str:
    try:
        sender = str(getattr(mail, "SenderEmailAddress", "") or "").strip()
        if sender and not sender.startswith("/"):
            return sender.lower()
    except Exception:
        pass

    try:
        sender_obj = mail.Sender
        exchange_user = sender_obj.GetExchangeUser() if sender_obj else None
        if exchange_user and exchange_user.PrimarySmtpAddress:
            return str(exchange_user.PrimarySmtpAddress).strip().lower()
    except Exception:
        pass

    return ""


def get_message_id(mail) -> str:
    try:
        return str(
            mail.PropertyAccessor.GetProperty(
                "http://schemas.microsoft.com/mapi/proptag/0x1035001E"
            )
            or ""
        ).strip()
    except Exception:
        pass

    try:
        return str(getattr(mail, "EntryID", "") or "").strip()
    except Exception:
        return ""


def iter_outlook_folders(folder):
    yield folder
    if not OUTLOOK_SCAN_SUBFOLDERS:
        return

    try:
        count = int(folder.Folders.Count)
    except Exception:
        count = 0

    for idx in range(1, count + 1):
        try:
            child = folder.Folders.Item(idx)
            yield from iter_outlook_folders(child)
        except Exception:
            continue


def iter_outlook_messages_since(start: datetime, until: Optional[datetime] = None):
    namespace = get_outlook_namespace()
    inbox = namespace.GetDefaultFolder(6)
    seen_ids: set[str] = set()

    for folder in iter_outlook_folders(inbox):
        try:
            items = folder.Items
            items.Sort("[ReceivedTime]", True)
            count = int(items.Count)
        except Exception:
            continue

        for idx in range(1, count + 1):
            try:
                mail = items.Item(idx)
                received = normalize_datetime(getattr(mail, "ReceivedTime", None))
                if received is None:
                    continue
                if received < start:
                    break
                if until is not None and received > until:
                    continue

                message_id = get_message_id(mail)
                dedupe = message_id or str(getattr(mail, "EntryID", "") or "")
                if dedupe and dedupe in seen_ids:
                    continue
                if dedupe:
                    seen_ids.add(dedupe)

                yield mail, received
            except Exception:
                continue


def find_target_attachment(mail):
    target = normalize_filename(TARGET_ATTACHMENT)
    try:
        count = int(mail.Attachments.Count)
    except Exception:
        return None

    for idx in range(1, count + 1):
        try:
            attachment = mail.Attachments.Item(idx)
            if normalize_filename(str(attachment.FileName or "")) == target:
                return attachment
        except Exception:
            continue
    return None


def get_attachment_count(mail) -> int:
    try:
        return int(mail.Attachments.Count)
    except Exception:
        return 0


def ensure_log_workbook():
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)

    if LOG_PATH.exists():
        wb = load_workbook(LOG_PATH)
        ws = wb.active
        if ws.max_row < 1:
            ws.append(LOG_HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Processamento"
        ws.append(LOG_HEADERS)

    headers = [cell.value for cell in ws[1]]
    if headers != LOG_HEADERS:
        ws.delete_rows(1, ws.max_row)
        ws.append(LOG_HEADERS)

    return wb, ws


def save_log_workbook(wb) -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(LOG_PATH)


def existing_logged_message_ids(ws) -> set[str]:
    ids: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        message_id = str(row[4] or "").strip()
        result = normalize_text(str(row[14] or ""))
        completed = (
            result.startswith("ok")
            or result.startswith("historico")
            or result.startswith("sem anexo")
        )
        if message_id and completed:
            ids.add(message_id)
    return ids


def append_log_row(
    ws,
    *,
    received: datetime,
    sender: str,
    subject: str,
    message_id: str,
    parsed: Optional[dict],
    attachment_found: bool,
    destination: Optional[Path],
    mode: str,
    action: str,
    result: str,
    files_written: int = 0,
    missing_conventions: int = 0,
    cc_fallback: int = 0,
    notes: str = "",
) -> None:
    parsed = parsed or {}
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        received.strftime("%Y-%m-%d %H:%M:%S"),
        sender,
        subject,
        message_id,
        parsed.get("tipo", ""),
        parsed.get("area", ""),
        parsed.get("ano", ""),
        parsed.get("mes", ""),
        parsed.get("estado", ""),
        "Sim" if attachment_found else "Não",
        str(destination or ""),
        mode,
        action,
        result,
        files_written,
        missing_conventions,
        cc_fallback,
        notes,
    ])


def safe_write_bytes(path: Path, data: bytes) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        existing = path.read_bytes()
        if existing == data:
            return "JA_EXISTIA_IDENTICO"
        raise FileExistsError(f"Já existe um ficheiro diferente no destino: {path}")
    path.write_bytes(data)
    return "GRAVADO"


def save_attachment_to_bytes(attachment) -> bytes:
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir) / str(attachment.FileName or TARGET_ATTACHMENT)
        attachment.SaveAsFile(str(tmp_path))
        return tmp_path.read_bytes()


def sanitize_filename_component(value: str, max_len: int = 90) -> str:
    value = str(value or "").strip()
    value = re.sub(r'[<>:"/\\|?*]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value or "sem_assunto")[:max_len].rstrip(" .")


def unique_destination_path(path: Path, data: Optional[bytes] = None) -> Path:
    if not path.exists():
        return path
    if data is not None:
        try:
            if path.read_bytes() == data:
                return path
        except Exception:
            pass
    for idx in range(2, 1000):
        candidate = path.with_name(f"{path.stem}_{idx}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Não foi possível gerar nome único para: {path}")


def save_reclamacao_email(
    mail, received: datetime, destination: Path, message_id: str
) -> dict:
    destination.mkdir(parents=True, exist_ok=True)
    subject = str(getattr(mail, "Subject", "") or "").strip()
    digest_source = message_id or f"{received.isoformat()}|{subject}"
    digest = hashlib.sha1(
        digest_source.encode("utf-8", errors="replace")
    ).hexdigest()[:10]
    base = (
        f"{received:%Y%m%d_%H%M%S}_"
        f"{sanitize_filename_component(subject, 70)}_{digest}"
    )

    files_written = 0
    saved_files: list[str] = []
    notes: list[str] = []

    msg_path = destination / f"{base}.msg"
    if not msg_path.exists():
        try:
            mail.SaveAs(str(msg_path), 3)
            files_written += 1
            saved_files.append(msg_path.name)
        except Exception as exc:
            notes.append(f"Não foi possível guardar o .msg: {exc}")
    else:
        saved_files.append(msg_path.name)

    for idx in range(1, get_attachment_count(mail) + 1):
        try:
            attachment = mail.Attachments.Item(idx)
            filename = sanitize_filename_component(
                str(attachment.FileName or f"anexo_{idx}")
            )
            with tempfile.TemporaryDirectory() as tmp_dir:
                tmp_path = Path(tmp_dir) / filename
                attachment.SaveAsFile(str(tmp_path))
                data = tmp_path.read_bytes()

            target = unique_destination_path(destination / filename, data=data)
            if target.exists() and target.read_bytes() == data:
                saved_files.append(target.name)
                continue

            target.write_bytes(data)
            files_written += 1
            saved_files.append(target.name)
        except Exception as exc:
            notes.append(f"Anexo {idx}: {exc}")

    return {
        "files_written": files_written,
        "attachment_count": get_attachment_count(mail),
        "saved_files": saved_files,
        "action": "Reclamação guardada (.msg + anexos)",
        "notes": "; ".join(notes),
    }


def extract_all_files_from_zip_bytes(
    zip_bytes: bytes,
    origem: str,
    password: Optional[bytes] = None,
):
    found = []
    with zipfile.ZipFile(io.BytesIO(zip_bytes), "r") as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            basename = Path(info.filename.replace("\\", "/")).name
            if not basename:
                continue
            try:
                data = zf.read(info)
            except RuntimeError:
                data = zf.read(info, pwd=password) if password else zf.read(info)

            if basename.lower().endswith(".zip"):
                found.extend(
                    extract_all_files_from_zip_bytes(
                        data, f"{origem} > {info.filename}", password
                    )
                )
            else:
                found.append((basename, data, origem))
    return found


def process_attachment_for_area(
    zip_bytes: bytes,
    parsed: dict,
    destination: Path,
    process_zip_callback: Callable,
    mapping_dict: dict,
    zip_password: bytes,
):
    destination.mkdir(parents=True, exist_ok=True)
    files_written = 0
    notes: list[str] = []

    original_zip_path = destination / TARGET_ATTACHMENT
    zip_status = safe_write_bytes(original_zip_path, zip_bytes)
    if zip_status == "GRAVADO":
        files_written += 1
    else:
        notes.append("ZIP original já existia e era idêntico.")

    if parsed["area"] == "Saúde Oral":
        for filename, content, _origem in extract_all_files_from_zip_bytes(
            zip_bytes, TARGET_ATTACHMENT, zip_password
        ):
            if safe_write_bytes(destination / filename, content) == "GRAVADO":
                files_written += 1
        return {
            "files_written": files_written,
            "missing_conventions": 0,
            "cc_fallback": 0,
            "action": "Descompactado sem transformação",
            "notes": "; ".join(notes),
        }

    outputs, report = process_zip_callback(
        zip_bytes, TARGET_ATTACHMENT, mapping_dict
    )

    for item in outputs:
        if safe_write_bytes(
            destination / item["filename"], item["data"]
        ) == "GRAVADO":
            files_written += 1

    missing_total = sum(int(item.get("missing") or 0) for item in report)
    cc_fallback = sum(
        int((item.get("cc_stats") or {}).get("fallback") or 0)
        for item in report
    )

    if not outputs:
        notes.append("Não foram encontrados TXT para transformar.")

    return {
        "files_written": files_written,
        "missing_conventions": missing_total,
        "cc_fallback": cc_fallback,
        "action": "ZIP guardado + ficheiros transformados",
        "notes": "; ".join(notes),
    }


def build_historical_log_2026() -> dict:
    wb, ws = ensure_log_workbook()
    existing_ids = existing_logged_message_ids(ws)

    stats = {"encontrados": 0, "adicionados": 0, "ja_existiam": 0}

    for mail, received in iter_outlook_messages_since(HISTORY_FROM, HISTORY_UNTIL):
        sender = get_sender_smtp(mail)
        if normalize_text(sender) != normalize_text(CCMSNS_SENDER):
            continue

        subject = str(getattr(mail, "Subject", "") or "").strip()
        body = get_mail_body(mail)
        parsed = (
            parse_reclamacao_email(subject, body)
            if is_reclamacao_email(subject, body)
            else parse_ccmsns_subject(subject)
        )
        if not parsed:
            continue

        stats["encontrados"] += 1
        message_id = get_message_id(mail)
        if message_id and message_id in existing_ids:
            stats["ja_existiam"] += 1
            continue

        destination = build_destination(parsed)
        append_log_row(
            ws,
            received=received,
            sender=sender,
            subject=subject,
            message_id=message_id,
            parsed=parsed,
            attachment_found=get_attachment_count(mail) > 0,
            destination=destination,
            mode="HISTÓRICO",
            action="Registado como já tratado",
            result="HISTÓRICO / JÁ TRATADO",
            notes="Sem cópia nem transformação de ficheiros históricos.",
        )
        if message_id:
            existing_ids.add(message_id)
        stats["adicionados"] += 1

    save_log_workbook(wb)
    stats["log_path"] = str(LOG_PATH)
    return stats


def process_new_ccmsns_emails(
    process_zip_callback: Callable,
    mapping_dict: dict,
    zip_password: bytes,
):
    wb, ws = ensure_log_workbook()
    existing_ids = existing_logged_message_ids(ws)

    stats = {
        "encontrados": 0,
        "processados": 0,
        "reclamacoes": 0,
        "reclamacoes_classificacao": 0,
        "ja_registados": 0,
        "sem_anexo": 0,
        "assunto_invalido": 0,
        "erros": 0,
    }
    details: list[dict] = []

    for mail, received in iter_outlook_messages_since(PROCESS_FROM):
        sender = get_sender_smtp(mail)
        if normalize_text(sender) != normalize_text(CCMSNS_SENDER):
            continue

        subject = str(getattr(mail, "Subject", "") or "").strip()
        body = get_mail_body(mail)
        is_reclamacao = is_reclamacao_email(subject, body)
        is_despesa = parse_ccmsns_subject(subject) is not None

        if not is_despesa and not is_reclamacao:
            continue

        stats["encontrados"] += 1
        message_id = get_message_id(mail)
        if message_id and message_id in existing_ids:
            stats["ja_registados"] += 1
            continue

        if is_reclamacao:
            parsed = parse_reclamacao_email(subject, body)
            destination = build_destination(parsed or {})
            missing_parts = []
            if not parsed or not parsed.get("area"):
                missing_parts.append("tipo/área")
            if not parsed or not parsed.get("mes"):
                missing_parts.append("mês")
            if not parsed or not parsed.get("ano"):
                missing_parts.append("ano")

            if destination is None:
                stats["reclamacoes_classificacao"] += 1
                notes = "Falta determinar: " + ", ".join(missing_parts)
                append_log_row(
                    ws,
                    received=received,
                    sender=sender,
                    subject=subject,
                    message_id=message_id,
                    parsed=parsed,
                    attachment_found=get_attachment_count(mail) > 0,
                    destination=None,
                    mode="AUTOMÁTICO",
                    action="Reclamação detetada — classificação necessária",
                    result="RECLAMAÇÃO/CLASSIFICAÇÃO NECESSÁRIA",
                    notes=notes,
                )
                details.append({
                    "Data": received.strftime("%Y-%m-%d %H:%M"),
                    "Assunto": subject,
                    "Tipo": "Reclamação",
                    "Resultado": notes,
                })
                continue

            try:
                proc = save_reclamacao_email(
                    mail, received, destination, message_id
                )
                result = "OK COM AVISOS" if proc["notes"] else "OK"
                append_log_row(
                    ws,
                    received=received,
                    sender=sender,
                    subject=subject,
                    message_id=message_id,
                    parsed=parsed,
                    attachment_found=proc["attachment_count"] > 0,
                    destination=destination,
                    mode="AUTOMÁTICO",
                    action=proc["action"],
                    result=result,
                    files_written=proc["files_written"],
                    notes=proc["notes"],
                )
                if message_id:
                    existing_ids.add(message_id)
                stats["processados"] += 1
                stats["reclamacoes"] += 1
                details.append({
                    "Data": received.strftime("%Y-%m-%d %H:%M"),
                    "Assunto": subject,
                    "Tipo": "Reclamação",
                    "Destino": str(destination),
                    "Resultado": result,
                })
            except Exception as exc:
                stats["erros"] += 1
                append_log_row(
                    ws,
                    received=received,
                    sender=sender,
                    subject=subject,
                    message_id=message_id,
                    parsed=parsed,
                    attachment_found=get_attachment_count(mail) > 0,
                    destination=destination,
                    mode="AUTOMÁTICO",
                    action="Erro durante processamento da reclamação",
                    result="ERRO",
                    notes=str(exc),
                )
                details.append({
                    "Data": received.strftime("%Y-%m-%d %H:%M"),
                    "Assunto": subject,
                    "Tipo": "Reclamação",
                    "Resultado": f"ERRO: {exc}",
                })
            continue

        parsed = parse_ccmsns_subject(subject)
        if not parsed or not parsed.get("area"):
            stats["assunto_invalido"] += 1
            append_log_row(
                ws,
                received=received,
                sender=sender,
                subject=subject,
                message_id=message_id,
                parsed=parsed,
                attachment_found=False,
                destination=None,
                mode="AUTOMÁTICO",
                action="Não processado",
                result="ASSUNTO/ÁREA NÃO RECONHECIDO",
                notes="Formato de assunto ou área não reconhecido.",
            )
            details.append({
                "Data": received.strftime("%Y-%m-%d %H:%M"),
                "Assunto": subject,
                "Tipo": "Despesa",
                "Resultado": "Assunto/área não reconhecido",
            })
            continue

        destination = build_destination(parsed)
        attachment = find_target_attachment(mail)

        if attachment is None:
            stats["sem_anexo"] += 1
            append_log_row(
                ws,
                received=received,
                sender=sender,
                subject=subject,
                message_id=message_id,
                parsed=parsed,
                attachment_found=False,
                destination=destination,
                mode="AUTOMÁTICO",
                action="Não processado",
                result="SEM ANEXO ULS LITORAL ALENTEJANO",
                notes=f"Não foi encontrado o anexo {TARGET_ATTACHMENT}.",
            )
            if message_id:
                existing_ids.add(message_id)
            details.append({
                "Data": received.strftime("%Y-%m-%d %H:%M"),
                "Assunto": subject,
                "Tipo": "Despesa",
                "Resultado": "Sem anexo ULS Litoral Alentejano",
            })
            continue

        try:
            zip_bytes = save_attachment_to_bytes(attachment)
            proc = process_attachment_for_area(
                zip_bytes,
                parsed,
                destination,
                process_zip_callback,
                mapping_dict,
                zip_password,
            )
            result = (
                "OK COM AVISOS"
                if proc["missing_conventions"] > 0 or proc["cc_fallback"] > 0
                else "OK"
            )
            append_log_row(
                ws,
                received=received,
                sender=sender,
                subject=subject,
                message_id=message_id,
                parsed=parsed,
                attachment_found=True,
                destination=destination,
                mode="AUTOMÁTICO",
                action=proc["action"],
                result=result,
                files_written=proc["files_written"],
                missing_conventions=proc["missing_conventions"],
                cc_fallback=proc["cc_fallback"],
                notes=proc["notes"],
            )
            if message_id:
                existing_ids.add(message_id)
            stats["processados"] += 1
            details.append({
                "Data": received.strftime("%Y-%m-%d %H:%M"),
                "Assunto": subject,
                "Tipo": "Despesa",
                "Destino": str(destination),
                "Resultado": result,
            })
        except Exception as exc:
            stats["erros"] += 1
            append_log_row(
                ws,
                received=received,
                sender=sender,
                subject=subject,
                message_id=message_id,
                parsed=parsed,
                attachment_found=True,
                destination=destination,
                mode="AUTOMÁTICO",
                action="Erro durante processamento",
                result="ERRO",
                notes=str(exc),
            )
            details.append({
                "Data": received.strftime("%Y-%m-%d %H:%M"),
                "Assunto": subject,
                "Tipo": "Despesa",
                "Destino": str(destination),
                "Resultado": f"ERRO: {exc}",
            })

    save_log_workbook(wb)
    return stats, details


def render_ccmsns_section(
    process_zip_callback: Callable,
    mapping_dict: dict,
    zip_password: bytes,
) -> None:
    st.divider()
    st.header("📧 Emails CCMSNS")
    st.write(f"**Remetente:** {CCMSNS_SENDER}")
    st.write(
        f"**Despesa faturada/conferida — anexo procurado:** {TARGET_ATTACHMENT}"
    )
    st.write(
        "**Reclamações:** o email completo (.msg) e os anexos são guardados "
        "em Ano\\Mês\\Tipo\\Reclamações."
    )
    st.write(
        "**Processamento automático apenas para emails recebidos "
        "a partir de 16/09/2026.**"
    )
    st.write(f"**Log:** {LOG_PATH}")

    local_windows = os.name == "nt"
    if not local_windows:
        st.info(
            "A automatização Outlook/CCMSNS só está disponível na execução "
            "local em Windows. O processamento manual de ZIP/TXT continua disponível."
        )
    elif not CONFERENCIA_ROOT.exists():
        st.warning(
            "A pasta de rede de Conferência não está acessível neste momento: "
            f"{CONFERENCIA_ROOT}"
        )

    if st.button(
        "🔌 Testar ligação ao Outlook",
        disabled=not local_windows,
    ):
        test = testar_ligacao_outlook()
        if test["ok"]:
            st.success("✅ " + test["mensagem"])
        else:
            st.error("❌ " + test["mensagem"])

    col_hist, col_auto = st.columns(2)

    with col_hist:
        st.subheader("Histórico 2026")
        st.caption(
            "Pesquisa despesas, conferências e reclamações de 01/01/2026 a "
            "15/09/2026 e acrescenta-as ao Excel como já tratadas, sem copiar "
            "nem transformar ficheiros históricos."
        )
        if st.button(
            "🕘 Construir/atualizar log histórico",
            use_container_width=True,
            disabled=not local_windows,
        ):
            try:
                with st.spinner("A pesquisar emails históricos no Outlook..."):
                    hist = build_historical_log_2026()
                st.success(
                    "✅ Log histórico atualizado. "
                    f"Encontrados: {hist['encontrados']} | "
                    f"Adicionados: {hist['adicionados']} | "
                    f"Já existentes: {hist['ja_existiam']}"
                )
                st.code(hist["log_path"])
            except Exception as exc:
                st.error(f"❌ Erro no histórico: {exc}")

    with col_auto:
        st.subheader("Novos emails")
        st.caption(
            "Procura despesas/conferências e reclamações desde 16/09/2026. "
            "Só processa mensagens ainda não concluídas no log."
        )
        if st.button(
            "🚀 Processar novos emails CCMSNS",
            type="primary",
            use_container_width=True,
            disabled=not local_windows,
        ):
            try:
                with st.spinner("A pesquisar e processar novos emails..."):
                    stats, details = process_new_ccmsns_emails(
                        process_zip_callback,
                        mapping_dict,
                        zip_password,
                    )
                st.success(
                    "✅ Pesquisa concluída. "
                    f"Encontrados: {stats['encontrados']} | "
                    f"Processados: {stats['processados']} | "
                    f"Reclamações: {stats['reclamacoes']} | "
                    f"Reclamações por classificar: "
                    f"{stats['reclamacoes_classificacao']} | "
                    f"Já registados: {stats['ja_registados']} | "
                    f"Sem anexo: {stats['sem_anexo']} | "
                    f"Erros: {stats['erros']}"
                )
                if details:
                    st.dataframe(
                        pd.DataFrame(details),
                        use_container_width=True,
                        hide_index=True,
                    )
            except Exception as exc:
                st.error(f"❌ Erro no processamento de emails: {exc}")

    st.info(
        "Saúde Oral/SISO: o ZIP é guardado e descompactado com a password "
        "habitual, mas os ficheiros não passam pela transformação "
        "MCDT/Termas/Centros de Custo."
    )
    st.info(
        "Reclamações: quando área, mês e ano são identificados, são guardadas "
        "em Ano\\Mês\\Tipo\\Reclamações. Se faltar algum desses elementos, "
        "o email é sinalizado para classificação e não é marcado como concluído."
    )
