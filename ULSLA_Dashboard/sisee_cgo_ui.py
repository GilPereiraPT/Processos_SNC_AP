"""Página CGO integrada no módulo Produção de Mapas/SISEE."""
from pathlib import Path
from decimal import Decimal
from io import BytesIO
import hashlib,json
import streamlit as st
from ulsla_branding import period_selectbox
from openpyxl import load_workbook
from sisee_cgo import (load_inputs,save_inputs,validate_inputs,read_imported_ba,
 read_calculation_xls,calculate,fill_template,MANUAL_ROWS,VERSAO_CGO)


def render_cgo():
    st.subheader('Controlo de gastos operacionais')
    import sisee_cgo
    st.caption(f'CGO {VERSAO_CGO} · Motor ativo: {Path(sisee_cgo.__file__).resolve()}')
    if st.session_state.get('cgo_motor_version') != VERSAO_CGO:
        st.session_state.pop('cgo_result',None)
        st.session_state['cgo_motor_version']=VERSAO_CGO
    a,b=st.columns(2)
    year=period_selectbox('Exercício',list(range(2024,2031)),index=2,key='cgo_year', target=a)
    month=period_selectbox('Trimestre',[3,6,9,12],index=2,format_func=lambda m:f'T{m//3} · acumulado até {m:02d}',key='cgo_month', target=b)
    period=f'{year}-T{month//3}'
    try:defaults,origin=load_inputs(period)
    except (ValueError,OSError,json.JSONDecodeError) as exc:
        st.error('Não foi possível ler os dados guardados: '+str(exc));return
    st.caption('Preenchimento inicial: '+origin+'. Confirma ou altera os valores para o período escolhido.')
    labels={'doentes':'Número de doentes / denominador do indicador','rh_inicio':'RH no início do ano',
      'entradas':'Entradas desde o início do ano','saidas':'Saídas desde o início do ano',
      'rh_medio':'Número médio de RH','orgaos':'Número de órgãos sociais','direcao':'Número de cargos de direção',
      'viaturas':'Número de viaturas','valorizacoes':'Acréscimo por valorizações remuneratórias (€)',
      'novos_trabalhadores':'Gastos pelo aumento do número de trabalhadores (€)'}
    values={}
    with st.form('cgo_complementares_'+period):
        columns=st.columns(3)
        for i,(key,label) in enumerate(labels.items()):
            integer=key not in ('rh_medio','valorizacoes','novos_trabalhadores')
            initial=int(defaults[key]) if integer else float(defaults[key])
            values[key]=columns[i%3].number_input(label,min_value=0 if integer else 0.0,value=initial,
              step=1 if integer else 0.01,key='cgo_'+period+'_'+key)
        st.caption('O denominador começa em 106 110 e pode ser alterado. O número de RH no final e o número de trabalhadores são calculados a partir dos dados acima.')
        with st.expander('Rubricas sem correspondência contabilística específica — confirmar valores'):
            st.caption('Valores iniciais zero, conforme o Excel fornecido. Os valores guardados são reutilizados; atualiza se existir despesa nestas rubricas. Não são inferidos apenas por estarem ausentes no balancete.')
            for row,label in MANUAL_ROWS.items():
                key=f'linha_{row}'
                values[key]=st.number_input(label+' (€)',min_value=0.0,value=float(defaults[key]),step=.01,key='cgo_'+period+'_'+key)
        saved=st.form_submit_button('Guardar dados complementares',type='primary')
    valid=True
    try:validate_inputs(values)
    except ValueError as exc:st.error(str(exc));valid=False
    if saved and valid:
        try:save_inputs(period,values);st.success('Dados guardados para este período e como valores por defeito dos próximos formulários CGO.')
        except OSError as exc:st.error('Não foi possível guardar os dados: '+str(exc));return
    st.markdown('#### Dados contabilísticos e modelo oficial')
    source=st.radio('Origem do balancete',['BA já importado no Dashboard','Excel de cálculo com folha Balancete (.xls)'],key='cgo_source')
    balance_file=None
    if source.startswith('Excel'):
        balance_file=st.file_uploader('Carregar Excel de cálculo / balancete',type=['xls'],key='cgo_balance')
        source_confirmed=st.checkbox('Confirmo que os valores do balancete correspondem ao acumulado até ao trimestre selecionado.',key='cgo_period_confirmation_'+period)
        st.caption('O ficheiro de exemplo mantém cabeçalhos antigos. A data do balancete deve ser confirmada pelo utilizador.')
    else:source_confirmed=True
    template=st.file_uploader('Modelo oficial CGO Setor Saúde (.xlsx)',type=['xlsx'],key='cgo_template')
    absent=st.checkbox('Confirmo que as subcontas ausentes no balancete completo não têm movimento/saldo a reportar.',key='cgo_absent_'+period)
    fingerprint=hashlib.sha256((json.dumps(values,sort_keys=True)+period+source+str(source_confirmed)+str(absent)).encode()
        +(template.getvalue() if template else b'')+(balance_file.getvalue() if balance_file else b'')).hexdigest()
    # Alterar campos/período/modelo invalida o ficheiro anterior apresentado.
    old=st.session_state.get('cgo_result')
    if old and old['fingerprint']!=fingerprint:st.session_state.pop('cgo_result',None)
    if st.button('Preencher formulário CGO',type='primary',disabled=not valid):
        st.session_state.pop('cgo_result',None)
        try:
            if template is None:raise ValueError('Carrega o modelo oficial CGO.')
            if not source_confirmed:raise ValueError('Confirma o período do balancete carregado.')
            accounts=(read_calculation_xls(balance_file.getvalue()) if balance_file else None) if source.startswith('Excel') else read_imported_ba(year,month)
            if accounts is None:raise ValueError('Carrega o Excel com a folha Balancete.')
            rows,warnings,trace=calculate(accounts,values)
            if warnings and not absent:raise ValueError(warnings[0]+'. Confirma a ausência de movimentos ou completa o balancete.')
            binary,count=fill_template(template.getvalue(),year,month,rows)
            save_inputs(period,values)
            st.session_state['cgo_result']={'fingerprint':fingerprint,'binary':binary,'filename':template.name,'rows':rows,'count':count,'warnings':warnings}
        except Exception as exc:st.error('Não foi possível preencher o CGO: '+str(exc))
    result=st.session_state.get('cgo_result')
    if result:
        st.success(f"{result['count']} células preenchidas em {period}. Outros períodos e cabeçalhos preservados.")
        rows=result['rows'];a,b,c=st.columns(3)
        def euro(n):return f'{n:,.2f} €'.replace(',','X').replace('.',',').replace('X','.')
        a.metric('Gastos operacionais',euro(rows[77]));b.metric('Gastos por doente',euro(rows[98]));c.metric('Número de doentes',f"{values['doentes']:,}".replace(',',' '))
        for warning in result['warnings']:st.warning(warning+' — ausência confirmada pelo utilizador.')
        st.download_button('Descarregar formulário CGO preenchido',result['binary'],file_name=result['filename'],mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        with st.expander('Conferir valores e regras aplicadas'):
            wb=load_workbook(BytesIO(result['binary']),read_only=True,data_only=True);sh=wb['Formulário']
            st.dataframe([{'Linha':r,'Rubrica':sh.cell(r,2).value,'Valor':float(v)} for r,v in sorted(rows.items())],hide_index=True,use_container_width=True)
            wb.close()
            st.caption('Consultoria excluída de Outros serviços especializados; estudos/pareceres/consultadoria = consultoria + auditoria. Gastos com trabalhadores = total pessoal − total órgãos sociais − direção. Remunerações fixas excluem subsídio de refeição para evitar duplicação. Gastos líquidos = débito − crédito; rendimentos = crédito − débito. Os resíduos e os subtotais são reconciliados antes de exportar.')
