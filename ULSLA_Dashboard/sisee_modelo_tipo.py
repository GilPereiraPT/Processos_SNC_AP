"""Identificação do formulário SISEE pelos identificadores oficiais presentes no Excel.
Não depende do nome do ficheiro e nunca altera o documento.
"""
from io import BytesIO
from openpyxl import load_workbook

def identificar_modelo(xlsx_bytes):
    wb=load_workbook(BytesIO(xlsx_bytes),read_only=True,data_only=False)
    try:
        if 'Formulário' not in wb.sheetnames:
            raise ValueError('Folha Formulário não encontrada')
        sh=wb['Formulário']
        b8=str(sh['B8'].value or '').upper()
        b14=str(sh['B14'].value or '').casefold()
        b15=str(sh['B15'].value or '').casefold()
        a14=str(sh['A14'].value or '').strip()
        a15=str(sh['A15'].value or '').strip()
        if str(sh['A8'].value).strip()=='1759' and 'DEMONSTRAÇÃO DOS RESULTADOS' in b8:
            return 'DR'
        if str(sh['A8'].value).strip()=='1752' and b8=='BALANÇO':
            return 'BAL'
        if str(sh['A8'].value).strip()=='2069' and 'TESOURARIA' in str(sh['D8'].value).upper():
            return 'UTE'
        if str(sh['A8'].value).strip()=='1629' and 'CONTROLO DE GASTOS OPERACIONAIS' in b8:
            return 'CGO'
        if a14=='663' and a15=='664' and 'aquisições de bens e serviços' in b14 and 'aquisições de capital' in b15:
            return 'SD'
        if 'PRAZO MÉDIO DE PAGAMENTOS' in b8 and str(sh['A10'].value or '').strip()=='623':
            return 'PMP'
        if 'DEMONSTRAÇÃO DOS FLUXOS DE CAIXA' in b8 and str(sh['A10'].value or '').strip()=='1546':
            return 'DFCx'
        raise ValueError('Não corresponde aos formulários SD, PMP, DFCx, CGO, DR, BAL ou UTE conhecidos')
    finally:
        wb.close()

def exigir_modelo_correto(xlsx_bytes, selecionado):
    modelo=identificar_modelo(xlsx_bytes)
    esperado='SD' if str(selecionado).startswith('SD') else str(selecionado)
    if modelo!=esperado:
        raise ValueError(f'FORMULÁRIO INCORRETO: carregaste o modelo {modelo}, mas o menu está em {esperado}. Seleciona {modelo} no menu Formulário antes de gerar. Nada foi alterado.')
    return modelo
