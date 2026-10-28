"""Controlo de gastos operacionais SISEE: cálculos e dados complementares."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import json, os, tempfile
from io import BytesIO
from datetime import datetime, timezone
from openpyxl import load_workbook

VERSAO_CGO='2.5.2'
ROOT=Path(__file__).resolve().parent
STORE=ROOT/'dados_sisee'/'cgo_dados_complementares.json'
D=Decimal
ZERO=D('0')
DEFAULTS={'doentes':106110,'rh_inicio':1139,'entradas':25,'saidas':21,'rh_medio':1140,
 'orgaos':5,'direcao':4,'viaturas':44,'valorizacoes':0.0,'novos_trabalhadores':0.0}
MANUAL_ROWS={29:'Deslocações',40:'Comissões',51:'Órgãos sociais — prémios de gestão',
 52:'Órgãos sociais — outras remunerações',54:'Órgãos sociais — formação',
 55:'Órgãos sociais — pensões e planos de saúde',56:'Órgãos sociais — encargos sobre remunerações',
 57:'Órgãos sociais — acidentes no trabalho',58:'Órgãos sociais — seguros de saúde',
 59:'Órgãos sociais — seguros de vida',60:'Órgãos sociais — indemnizações',
 72:'Pessoal — seguros de saúde',73:'Pessoal — seguros de vida',95:'Indemnizações compensatórias'}
for r in MANUAL_ROWS:DEFAULTS[f'linha_{r}']=0.0


def read_store(path=STORE):
    if not path.exists():return {'periodos':{},'ultimos':{}}
    data=json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data.get('periodos'),dict) or not isinstance(data.get('ultimos'),dict):
        raise ValueError('Ficheiro de dados complementares inválido; verificar dados_sisee.')
    return data


def load_inputs(period,path=STORE):
    data=read_store(path)
    stored=data['periodos'].get(period)
    source='dados guardados para este período' if stored is not None else 'últimos dados guardados' if data['ultimos'] else 'valores iniciais do Excel fornecido'
    values={**DEFAULTS,**(stored if stored is not None else data['ultimos'])}
    return values,source


def validate_inputs(values):
    for k in DEFAULTS:
        if k not in values:raise ValueError('Dado complementar em falta: '+k)
        n=D(str(values[k]))
        if not n.is_finite() or n<0:raise ValueError('Valor inválido: '+k)
    for k in ['doentes','rh_inicio','entradas','saidas','orgaos','direcao','viaturas']:
        if D(str(values[k]))!=D(str(values[k])).to_integral_value():raise ValueError('É necessário um número inteiro: '+k)
    if values['doentes']<=0:raise ValueError('O número de doentes deve ser superior a zero.')
    end=values['rh_inicio']+values['entradas']-values['saidas']
    if end<values['orgaos']+values['direcao']:raise ValueError('Efetivos insuficientes para os órgãos sociais e cargos de direção indicados.')


def save_inputs(period,values,path=STORE):
    validate_inputs(values)
    data=read_store(path);data['periodos'][period]=dict(values);data['ultimos']=dict(values)
    data['ultimo_periodo']=period;data['guardado_em']=datetime.now(timezone.utc).isoformat()
    path.parent.mkdir(parents=True,exist_ok=True)
    fd,temp=tempfile.mkstemp(prefix='cgo_',suffix='.tmp',dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:json.dump(data,f,ensure_ascii=False,indent=2)
        os.replace(temp,path)
    finally:
        if os.path.exists(temp):os.unlink(temp)


def read_calculation_xls(content):
    try:import xlrd
    except ImportError as exc:raise ValueError('Para ler o Excel .xls é necessário xlrd. Usa o BA já importado ou instala xlrd no ambiente da aplicação.') from exc
    w=xlrd.open_workbook(file_contents=content)
    if 'Balancete' not in w.sheet_names():raise ValueError('O Excel deve conter a folha Balancete com a estrutura do ficheiro de cálculo fornecido.')
    sh=w.sheet_by_name('Balancete');out={}
    for i in range(9,sh.nrows):
        code=sh.cell_value(i,0)
        if code=='':continue
        code=str(int(code)) if isinstance(code,(int,float)) else str(code).strip().replace('.','')
        if not code.isdigit():continue
        if code in out:raise ValueError('Conta repetida no balancete: '+code)
        # Gastos: débito-crédito; rendimentos: crédito-débito.
        deb=D(str(sh.cell_value(i,4) or 0));cred=D(str(sh.cell_value(i,5) or 0))
        out[code]=cred-deb if code.startswith('7') else deb-cred
    return out


def read_imported_ba(year,month):
    from sisee_generator import period_obj
    obj=period_obj(year,month,'BA')
    if obj is None:raise ValueError(f'BA de {year}-{month:02d} não importado.')
    out={};seen_ids=set()
    for row in obj.get('registos',[]):
        code=''.join(x for x in str(row.get('ContaLocal','')) if x.isdigit())
        # Âmbito CGO; contas orçamentais/controlo são irrelevantes.
        if not code.startswith(('6','71','72')):
            continue
        vals=row.get('detalhes',{}).get('BA04')
        if vals is None:continue
        # As contas financeiras podem repartir-se por NIF e outras dimensões.
        # 999999990 não é um total que substitua as entidades parceiras.
        record_id=str(row.get('ID_Registo','')).strip()
        if record_id:
            if record_id in seen_ids:
                raise ValueError('ID_Registo financeiro repetido no BA: '+record_id)
            seen_ids.add(record_id)
        debit=D(str(vals.get('Debito',0)));credit=D(str(vals.get('Credito',0)))
        value=credit-debit if code.startswith('7') else debit-credit
        out[code]=out.get(code,ZERO)+value

    return out


def calculate(accounts,inputs):
    validate_inputs(inputs);missing=set();trace={}
    def balance(prefix):
        if prefix in accounts:return D(str(accounts[prefix]))
        children=sorted(k for k in accounts if k.startswith(prefix))
        # Fronteira sem sobreposição: não somar pais e filhos do balancete.
        frontier=[k for k in children if not any(k.startswith(p) and k!=p for p in children)]
        if not frontier:missing.add(prefix);return ZERO
        return sum((D(str(accounts[k])) for k in frontier),ZERO)
    for prefix in ('61','62','63','72'):
        if not any(k.startswith(prefix) for k in accounts):raise ValueError('Balancete sem a conta obrigatória '+prefix)
    rows={r:D(str(inputs[f'linha_{r}'])) for r in MANUAL_ROWS}
    direct={10:'611',13:'612411',14:'612412',15:'612419',16:'61242',17:'6126',18:'61243',19:'61244',20:'61245',
      23:'621',25:'62211',26:'62214',30:'6251',31:'6241',32:'6242',33:'6243',34:'6234',35:'6261',36:'6266',37:'6262',38:'6263',
      41:'6224',42:'6226',43:'6222',44:'6267',45:'6223',50:'63115',53:'63123',64:'63215',65:'6322051',
      67:'632203',68:'632208',69:'633',70:'635',71:'636',74:'634',76:'637',100:'62111'}
    for r,code in direct.items():rows[r]=balance(code);trace[r]=code
    rows[9]=balance('61');rows[11]=rows[9]-rows[10];rows[12]=sum(rows[r] for r in (13,14,15))
    rows[21]=rows[11]-sum(rows[r] for r in [12,16,17,18,19,20])
    rows[22]=balance('62');rows[27]=balance('6221')-rows[25]-rows[26];rows[24]=rows[25]+rows[26]+rows[27]
    rows[28]=rows[29]+rows[30];rows[39]=balance('625')-rows[28]
    parts=[23,24,28]+list(range(31,46));rows[46]=rows[22]-sum(rows[r] for r in parts)
    rows[47]=balance('63');rows[48]=balance('631');rows[49]=balance('6311')-rows[50]
    rows[61]=rows[48]-sum(rows[r] for r in range(49,61));rows[62]=rows[47]-rows[48]
    rows[63]=balance('6321')-rows[64];rows[66]=balance('6322')-rows[65]-rows[67]-rows[68]
    rows[75]=rows[62]-sum(rows[r] for r in list(range(63,75))+[76])
    rows[77]=rows[9]+rows[22]+rows[47]
    rows[78]=D(str(inputs['rh_inicio']));rows[79]=D(str(inputs['entradas']));rows[80]=D(str(inputs['saidas']))
    rows[81]=rows[78]+rows[79]-rows[80];rows[82]=D(str(inputs['rh_medio']))
    rows[83]=D(str(inputs['orgaos']));rows[84]=D(str(inputs['direcao']));rows[85]=rows[81]-rows[83]-rows[84];rows[86]=rows[83]+rows[84]+rows[85]
    rows[87]=sum(balance(k) for k in ['632117','632121','632147']);rows[88]=rows[47]-rows[48]-rows[87]
    rows[89]=D(str(inputs['viaturas']));rows[91]=rows[74];rows[92]=D(str(inputs['valorizacoes']));rows[93]=D(str(inputs['novos_trabalhadores']))
    rows[94]=balance('71')+balance('72');rows[96]=sum(balance(k) for k in ['626112','62631','6242','622611'])
    rows[97]=rows[25]+rows[26];rows[98]=rows[77]/D(str(inputs['doentes']))
    rows[99]=sum(balance(k) for k in ['6322041','6221911','622491'])
    for r in (21,27,39,46,61,75):
        if rows[r]<D('-.005'):raise ValueError(f'Residual negativo na linha {r}: {rows[r]:.2f}. Rever contas e dados complementares.')
    warnings=['Contas ausentes (não confundir com saldo zero): '+', '.join(sorted(missing))] if missing else []
    rounded={r:v.quantize(D('.01'),rounding=ROUND_HALF_UP) for r,v in rows.items()}
    # Controlos sobre os valores que serão gravados.
    controls=[(9,[10,11]),(11,list(range(12,22))),(12,[13,14,15]),(24,[25,26,27]),(28,[29,30]),
              (22,parts+[46]),(48,list(range(49,62))),(62,list(range(63,77))), (47,[48,62]),(77,[9,22,47]),(86,[83,84,85])]
    # No consumo, 12 já inclui 13–15; não voltar a somá-los ao total 11.
    controls[1]=(11,[12,16,17,18,19,20,21])
    for total,children in controls:
        if abs(rounded[total]-sum(rounded[r] for r in children))>D('.01'):
            raise ValueError('Total não reconciliado na linha '+str(total))
    return rounded,warnings,trace


def fill_template(content,year,month,values):
    if month not in (3,6,9,12):raise ValueError('Seleciona o mês de fecho do trimestre: 3, 6, 9 ou 12.')
    w=load_workbook(BytesIO(content));sh=w['Formulário']
    if str(sh['A8'].value)!='1629' or 'CONTROLO DE GASTOS' not in str(sh['B8'].value).upper():raise ValueError('Modelo incorreto: é necessário o formulário CGO Setor Saúde.')
    target=year*10000+(month//3)*10
    cols=[c.column for c in sh[6] if str(c.value)==str(target)]
    if len(cols)!=1:raise ValueError('Não foi encontrada uma única coluna para o trimestre selecionado.')
    col=cols[0]
    for row,value in values.items():
        cell=sh.cell(row,col)
        if cell.data_type=='f' or cell.__class__.__name__=='MergedCell':raise ValueError('Célula não editável: '+cell.coordinate)
        if not isinstance(sh.cell(row,1).value,(float,int)):raise ValueError('Linha sem identificador oficial: '+str(row))
        cell.value=float(value)
    out=BytesIO();w.save(out);return out.getvalue(),len(values)
