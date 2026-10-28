from pathlib import Path
import inspect
import hashlib
from io import BytesIO
from openpyxl import load_workbook
import streamlit as st
from ulsla_branding import period_selectbox
from ulsla_branding import hero
from sisee_mes import atualizar_mes
from sisee_divida import fill_sd, acss_excel
from sisee_modelo_tipo import identificar_modelo, exigir_modelo_correto

VERSAO_MAPAS = '2.6.0 · DR, Balanço, UTE e CGO'
st.set_page_config(page_title='ULSLA · Produção de Mapas', page_icon='🧾', layout='wide')
hero('ULSLA | Produção de Mapas',
     'Geração mensal de ficheiros para reporte institucional.',
     'Organizado por entidade de reporte: ACSS, SISEE e outros organismos.')

st.caption(f'**Produção de Mapas — {VERSAO_MAPAS}** · Motor ativo: `{Path(inspect.getfile(fill_sd)).resolve()}`')
if st.session_state.get('prod_mapas_build') != VERSAO_MAPAS:
    st.session_state.pop('prod_mapas_result',None)
    st.session_state.pop('sisee_sd_validacao',None)
    st.session_state['prod_mapas_build']=VERSAO_MAPAS

st.info('Os ficheiros PMP e DFCx usam valores acumulados desde janeiro. O SD usa o stock de dívida na data de referência. A antiguidade ACSS respeita os escalões próprios. Todos os ecrãs principais usam o logo da ULSLA, quando disponível localmente.')

ent_a, ent_b = st.columns([2,3])
with ent_a:
    entidade = st.selectbox('Entidade de reporte', ['SISEE', 'ACSS', 'Outros'])
with ent_b:
    if entidade == 'SISEE':
        mapa = st.selectbox('Mapa / formulário', ['PMP', 'DFCx', 'SD — Situação da Dívida', 'CGO — Gastos Operacionais', 'DR — Demonstração dos Resultados', 'BAL — Balanço', 'UTE — Unidade de Tesouraria do Estado'])
    elif entidade == 'ACSS':
        mapa = st.selectbox('Mapa / formulário', ['Antiguidade da Dívida'])
    else:
        mapa = st.selectbox('Mapa / formulário', ['Em preparação'])

if entidade == 'Outros':
    st.warning('Esta secção fica já preparada para futuros reportes a outros organismos, mantendo a mesma lógica de atualização do modelo original.')
    st.stop()

if entidade == 'SISEE' and mapa.split(' — ',1)[0] in ('DR','BAL','UTE'):
    import importlib.util
    import sys
    for module_name in ('sisee_modelo_tipo','sisee_financeiros','sisee_financeiros_ui'):
        module_path = Path(__file__).resolve().with_name(module_name + '.py')
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
    sys.modules['sisee_financeiros_ui'].render(mapa.split(' — ',1)[0])
    st.stop()

if entidade == 'SISEE' and mapa.startswith('CGO'):
    # Recarregar do diretório desta aplicação, evitando módulos antigos em memória.
    import importlib.util
    import sys
    for module_name in ('sisee_cgo', 'sisee_cgo_ui'):
        module_path = Path(__file__).resolve().with_name(module_name + '.py')
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
    sys.modules['sisee_cgo_ui'].render_cgo()
    st.stop()

with st.form('gerar_mapa'):
    a, b, c = st.columns(3)
    year = period_selectbox('Exercício', list(range(2024, 2031)), index=2, target=a)
    month = period_selectbox('Mês a atualizar', list(range(1, 13)), index=8, format_func=lambda x: f'{x:02d}', target=b)
    if entidade == 'SISEE' and mapa == 'PMP':
        origem_ativos = c.selectbox('Aquisições dos ativos (PMP)', ['Extratos contabilísticos', 'AFT XML'])
    else:
        origem_ativos = 'Extratos contabilísticos'
        c.markdown("<div style='padding-top:2rem;color:#6b7280'>Parâmetros específicos do mapa aplicados automaticamente.</div>", unsafe_allow_html=True)

    if entidade == 'SISEE':
        st.markdown('**Modelo oficial a atualizar**')
        template = st.file_uploader('Carregar o próprio ficheiro Excel (.xlsx). O mesmo nome será preservado no ficheiro devolvido.', type=['xlsx'])
        botao = st.form_submit_button('Atualizar o modelo carregado', type='primary')
    else:
        template = None
        botao = st.form_submit_button('Gerar mapa ACSS', type='primary')

if entidade == 'SISEE' and template is not None:
    try:
        identificado_visual = identificar_modelo(template.getvalue())
        pretendido = 'SD' if mapa.startswith('SD') else mapa
        if identificado_visual != pretendido:
            st.error(f'ATENÇÃO: carregaste um modelo {identificado_visual}, mas escolheste o processamento {pretendido}. O ficheiro não será alterado.')
        else:
            st.success(f'Modelo {identificado_visual} reconhecido. Pode ser atualizado.')
    except Exception as err:
        st.error(f'Modelo não reconhecido: {err}')

if botao:
    st.session_state.pop('prod_mapas_result', None)
    st.session_state.pop('sisee_sd_validacao', None)
    try:
        if entidade == 'SISEE' and template is not None:
            identificado = exigir_modelo_correto(template.getvalue(), mapa)
            st.caption(f'Modelo reconhecido: {identificado}. Processamento selecionado: {mapa}.')
        if entidade == 'SISEE' and template is None:
            st.error('Carrega primeiro o modelo Excel oficial a atualizar.')
        else:
            if entidade == 'ACSS':
                binary, meta = acss_excel(year, month)
                warnings = [
                    f"Origem: {meta['ficheiro']}",
                    f"Documentos: {meta['documentos']}",
                    f"Segurança Social reclassificada como Estado: {meta['seguranca_social_corrigida']} documentos",
                    f"Prazo de vencimento limitado a 60 dias: {meta['documentos_vencimento_ajustado']} documentos ({meta['valor_documentos_vencimento_ajustado']:,.2f} EUR)",
                ]
                filename = f'ULSLA_ACSS_Antiguidade_{year}_{month:02d}.xlsx'
                n = 1
            elif mapa.startswith('SD'):
                # Não aceitar modelos contaminados por uma exportação anterior.
                model_wb = load_workbook(BytesIO(template.getvalue()), read_only=True, data_only=False)
                sh_pre = model_wb['Formulário']
                extraneous = [f'K{r}' for r in range(10,38) if sh_pre[f'K{r}'].value is not None]
                if extraneous:
                    raise ValueError('O modelo carregado contém dados fora da tabela oficial ('+', '.join(extraneous[:5])+'). Usa o modelo original descarregado da SISEE.')
                binary, meta = fill_sd(template.getvalue(), year, month)
                # Verificação independente no ponto final antes do botão de descarregar.
                valid_wb=load_workbook(BytesIO(binary),read_only=True,data_only=False)
                valid=valid_wb['Formulário']
                cells=[valid.cell(r,c) for r in (14,15,16) for c in range(3,9)]+[valid['C17']]
                if len(cells)!=19 or any(not isinstance(c.value,(float,int)) for c in cells):
                    raise RuntimeError('Exportação SD bloqueada: células C14:H16 e C17 não preenchidas')
                if any(valid[f'K{r}'].value is not None for r in range(10,38)):
                    raise RuntimeError('Exportação SD bloqueada: foram detetados dados fora da tabela oficial')
                if abs(sum(valid.cell(16,c).value for c in range(3,9))-valid['C17'].value)>0.005:
                    raise RuntimeError('Exportação SD bloqueada: totais não reconciliam')
                warnings = [
                    f"Origem: {meta['ficheiro']}",
                    f"Documentos: {meta['documentos']}",
                    f"Segurança Social corrigida de Externos para Estado: {meta['seguranca_social_corrigida']} documentos",
                    f"Prazo de vencimento limitado a 60 dias: {meta['documentos_vencimento_ajustado']} documentos ({meta['valor_documentos_vencimento_ajustado']:,.2f} EUR)",
                ]
                filename = template.name
                n = 24
            else:
                binary, warnings, n = atualizar_mes(template.getvalue(), mapa, year, month, 'xml_aft' if origem_ativos == 'AFT XML' else 'extratos')
                filename = template.name

            # Evita resultados antigos quando se altera o mapa, mes ou ficheiro.
            if entidade == 'SISEE' and mapa.startswith('SD'):
                st.session_state['sisee_sd_validacao'] = {
                    'referencia': meta['data_referencia'],
                    'total_externos': meta['sisee_total_externos'],
                    'excluidos': meta['sisee_excluidos'],
                    'cf': meta.get('documentos_cf_excluidos',0),
                    'celulas': meta['sisee_celulas_preenchidas'],
                }
            else:
                st.session_state.pop('sisee_sd_validacao', None)
            if entidade == 'SISEE':
                tipo_saida = identificar_modelo(binary)
                tipo_esperado = 'SD' if mapa.startswith('SD') else mapa
                if tipo_saida != tipo_esperado:
                    raise RuntimeError(f'O ficheiro de saída é {tipo_saida}, mas era esperado {tipo_esperado}.')
            st.session_state['prod_mapas_result'] = (filename, binary, warnings, n, entidade, mapa, year, month)
    except Exception as ex:
        st.session_state.pop('prod_mapas_result', None)
        st.error(f'Não foi possível gerar o mapa: {ex}')

if 'prod_mapas_result' in st.session_state:
    filename, binary, warnings, n, entidade, mapa, year, month = st.session_state['prod_mapas_result']
    if entidade == 'ACSS':
        st.success(f'Mapa ACSS gerado para {month:02d}/{year}.')
    else:
        st.success(f'{n} célula(s) atualizada(s) em {month:02d}/{year}. O resto do modelo foi preservado.')
    if warnings:
        with st.expander(f'{len(warnings)} ponto(s) a validar antes da submissão'):
            for w in warnings:
                st.write('•', w)
    st.download_button(
        'Descarregar ficheiro',
        data=binary,
        file_name=filename,
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        type='primary',
    )

if 'sisee_sd_validacao' in st.session_state:
    info=st.session_state['sisee_sd_validacao']
    st.caption(f"SD validado: referência {info['referencia']} · Apenas fornecedores Externos · CF excluídos: {info['cf']}")
    st.metric('Total fornecedores Externos (SISEE)',f"{info['total_externos']:,.2f} €".replace(',','X').replace('.',',').replace('X','.'))
    with st.expander('Conferir valores efetivamente gravados no Excel (C14:H16 e C17)'):
        for row in (14,15,16,17):
            valores=[f'{cel}: {v:,.2f}'.replace(',','X').replace('.',',').replace('X','.') for cel,v in info['celulas'].items() if cel[1:].isdigit() and int(cel[1:])==row]
            st.write(' | '.join(valores))
