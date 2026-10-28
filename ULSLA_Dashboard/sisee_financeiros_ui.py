from pathlib import Path
from io import BytesIO
import hashlib,json,os,tempfile
import streamlit as st
from ulsla_branding import period_selectbox
from openpyxl import load_workbook
from sisee_financeiros import VERSION,ROOT,load_source,statement,ute,fill,check


def render(kind):
 titles={'DR':'Demonstração dos Resultados','BAL':'Balanço','UTE':'Unidade de Tesouraria do Estado'}
 st.subheader(titles[kind]);st.caption(f'SISEE {VERSION} · Motor ativo: {ROOT / "sisee_financeiros.py"}')
 a,b=st.columns(2)
 year=period_selectbox('Exercício',list(range(2024,2031)),index=2,key='fin_year_'+kind, target=a)
 month=period_selectbox('Trimestre',[3,6,9,12],index=2,format_func=lambda m:f'T{m//3} — fecho {m:02d}',key='fin_month_'+kind, target=b)
 period=f'{year}-T{month//3}';extras={}
 if kind=='UTE':
  st.info(f'Débitos e créditos exclusivamente de {month-2:02d} a {month:02d}/{year}. Saldos à data de início/fim do trimestre. Caixa separado dos bancos.')
  store=ROOT/'dados_sisee'/'ute_complementares.json'
  try:saved=json.loads(store.read_text(encoding='utf-8')) if store.exists() else {}
  except (ValueError,OSError) as exc:st.error('Erro na leitura dos dados UTE guardados: '+str(exc));return
  defaults=saved.get(period,{})
  with st.expander('Dados complementares UTE',expanded=True):
   st.caption('Valores guardados para este período, ou zero inicial. Estes campos não são deduzidos automaticamente dos movimentos bancários.')
   cols=st.columns(2)
   for i,(key,label) in enumerate([('juros_igcp','Ganhos financeiros do trimestre — IGCP'),('juros_outros','Ganhos financeiros do trimestre — Outros'),('dispensa','Montante autorizado de dispensa UTE'),('ganhos_dispensa','Ganhos financeiros correspondentes à dispensa'),('ganhos_entregues','Ganhos financeiros entregues ao Estado')]):
    extras[key]=cols[i%2].number_input(label+' (€)',min_value=0.0,value=float(defaults.get(key,0)),step=.01,key=f'fin_{period}_{key}')
   extras['autorizacao']=st.checkbox('Existe autorização de excecionamento UTE',value=bool(defaults.get('autorizacao',False)),key='fin_auth_'+period)
   st.caption('Os ganhos financeiros já incluídos nos débitos não voltam a ser adicionados ao saldo: o campo B é informativo.')
  with st.expander('Aplicações de tesouraria e financeiras — classificação por prazo'):
   st.caption('Preenche apenas se existirem aplicações. O saldo final é calculado por início − diminuições + aumentos. A classificação por prazo e os juros exigem confirmação.')
   for start,label in [(11,'Tesouraria IGCP'),(16,'Tesouraria Outros'),(22,'Financeiras IGCP'),(27,'Financeiras Outros')]:
    st.markdown('**'+label+'**');cols=st.columns(4)
    for i,(key,title) in enumerate([('inicio','Saldo inicial'),('ganhos','Ganhos do trimestre'),('diminuicoes','Diminuições'),('aumentos','Aumentos')]):
     field=f'ap_{start}_{key}'
     extras[field]=cols[i].number_input(title+' (€)',min_value=0.0,value=float(defaults.get(field,0)),step=.01,key=f'fin_{period}_{field}')
   extras['aplicacoes_confirmadas']=st.checkbox('Confirmo a classificação e os valores de todas as aplicações',value=False,key='fin_apps_confirm_'+period)
 else:st.caption('DR: acumulado desde janeiro, com os sinais do XML. Balanço: saldos à data do fecho. Só será atualizada a coluna do trimestre escolhido.')
 template=st.file_uploader('Carregar modelo oficial '+kind+' (.xlsx)',type=['xlsx'],key='fin_template_'+kind)
 if template:
  try:
   from sisee_modelo_tipo import identificar_modelo
   found=identificar_modelo(template.getvalue())
   if found!=kind:st.error('Modelo carregado '+found+'; selecionado '+kind+'.')
  except Exception as exc:st.error(str(exc))
 paths=sorted((ROOT/'dados_s3cp_xml').glob('*/*.json.gz')) if (ROOT/'dados_s3cp_xml').exists() else []
 signature=[(str(p),p.stat().st_mtime_ns,p.stat().st_size) for p in paths]
 fingerprint=hashlib.sha256((VERSION+kind+period+json.dumps(extras,sort_keys=True)+repr(signature)).encode()+(template.getvalue() if template else b'')).hexdigest()
 key='fin_result_'+kind
 if st.session_state.get(key,{}).get('fingerprint')!=fingerprint:st.session_state.pop(key,None)
 if st.button('Preencher '+kind,type='primary'):
  st.session_state.pop(key,None)
  try:
   if template is None:raise ValueError('Carrega o modelo oficial.')
   warnings=[]
   if kind=='UTE':values,warnings=ute(year,month,extras)
   else:
    values=statement(load_source(year,month,'DR' if kind=='DR' else 'BLC'),kind)
    counterpart='BLC' if kind=='DR' else 'DR'
    try:
     other=statement(load_source(year,month,counterpart),'BAL' if kind=='DR' else 'DR')
     check(values[37] if kind=='DR' else values[50],[other[50] if kind=='DR' else other[37]],'Resultado líquido DR = Balanço')
    except ValueError as exc:
     if 'não importado' in str(exc):warnings.append(str(exc)+' Não foi possível efetuar o confronto DR/Balanço.')
     else:raise
   binary,count=fill(template.getvalue(),kind,year,month,values)
   if kind=='UTE':
    saved[period]=extras;store.parent.mkdir(exist_ok=True)
    fd,tmp=tempfile.mkstemp(dir=store.parent,prefix='ute_',suffix='.tmp')
    try:
     with os.fdopen(fd,'w',encoding='utf-8') as h:json.dump(saved,h,ensure_ascii=False,indent=2)
     os.replace(tmp,store)
    finally:
     if os.path.exists(tmp):os.unlink(tmp)
   st.session_state[key]={'fingerprint':fingerprint,'binary':binary,'filename':template.name,'values':values,'count':count,'warnings':warnings}
  except Exception as exc:st.error('Não foi possível preencher '+kind+': '+str(exc))
 result=st.session_state.get(key)
 if result:
  st.success(f"{kind} {period}: {result['count']} células preenchidas. Cabeçalhos e outros períodos preservados.")
  for warning in result['warnings']:st.warning(warning)
  st.download_button('Descarregar '+kind+' preenchido',result['binary'],file_name=result['filename'],mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
  with st.expander('Conferir valores gravados',expanded=True):
   w=load_workbook(BytesIO(result['binary']),read_only=True);s=w['Formulário']
   st.dataframe([{'Linha':r,'Rubrica':s.cell(r,4 if kind=='UTE' else 2).value,'Valor':float(v)} for r,v in sorted(result['values'].items())],hide_index=True,use_container_width=True);w.close()
